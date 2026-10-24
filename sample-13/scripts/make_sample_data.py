"""サンプルの回答ファイル（3 教室ぶん）を作る。

    uv run python scripts/make_sample_data.py

教室ごとに、現場でよくある「形式のばらつき」をわざと入れている。
- みなと教室: Google フォームの回答 CSV（UTF-8・BOM 付き）。満足度は「5 (とても満足)」、きっかけは複数選択で「Instagram, チラシ」
- さくら教室: スタッフが手入力した Excel。上にタイトル行があり、見出しは 3 行目。年齢は数字、満足度に全角数字、再参加は ○ × △
- ひがし教室: 紙のアンケートを打ち込んだ CSV（Shift_JIS）。満足度は「満足」「やや不満」などの言葉、空行・同じ回答の二重入力・範囲外の値・電話番号入りのご意見が混ざる
回答の中身はすべて架空。乱数の種を固定しているので、何度実行しても同じファイルになる。
"""

from __future__ import annotations

import csv
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "sample_data"
rng = random.Random(20261024)

COMMENTS_GOOD = [
    "先生の説明がわかりやすく、家でも作れそうです。",
    "少人数で質問しやすかったです。",
    "生地をこねる感覚がつかめました。",
    "子どもと一緒に参加できてうれしかったです。",
    "レシピのカードがもらえて助かります。",
    "時間どおりに終わったので予定が立てやすかったです。",
    "道具がそろっていて手ぶらで来られました。",
    "次はパンのコースにも出てみたいです。",
]
COMMENTS_MID = [
    "もう少しゆっくり進めてもらえると助かります。",
    "駐車場の場所がわかりにくかったです。",
    "試食の時間がもう少しほしかったです。",
    "予約のメールが迷惑メールに入っていました。",
]
COMMENTS_LOW = [
    "人数が多くて先生に質問しにくかったです。",
    "開始が10分ほど遅れました。",
]
CHANNELS = ["Instagram", "チラシ", "知人の紹介", "Web検索", "その他"]


def pick_satisfaction() -> int:
    return rng.choices([5, 4, 3, 2, 1], weights=[40, 32, 15, 8, 5])[0]


def pick_comment(s: int) -> str:
    if rng.random() < 0.35:
        return ""
    pool = COMMENTS_GOOD if s >= 4 else COMMENTS_MID if s == 3 else COMMENTS_LOW
    return rng.choice(pool)


def pick_again(s: int) -> str:
    if s >= 4:
        return rng.choices(["はい", "未定"], weights=[85, 15])[0]
    if s == 3:
        return rng.choice(["はい", "未定", "いいえ"])
    return rng.choices(["いいえ", "未定"], weights=[60, 40])[0]


def pick_age() -> int:
    return rng.choices([17, 24, 33, 38, 45, 52, 61, 71], weights=[3, 14, 22, 18, 17, 12, 9, 5])[0] + rng.randint(0, 4)


def age_band(age: int) -> str:
    return "70代以上" if age >= 70 else f"{age // 10 * 10}代"


def lesson_days(n: int) -> list[date]:
    days = [date(2026, 10, 1) + timedelta(days=i) for i in range(31)]
    days = [d for d in days if d.weekday() in (2, 5, 6)]  # 水・土・日に開催
    return sorted(rng.choice(days) for _ in range(n))


# ----- みなと教室: Google フォーム -----
def make_minato() -> None:
    labels = {5: "5 (とても満足)", 4: "4", 3: "3", 2: "2", 1: "1 (不満)"}
    rows = []
    for d in lesson_days(46):
        s = pick_satisfaction()
        ts = datetime(d.year, d.month, d.day, rng.randint(11, 20), rng.randint(0, 59), rng.randint(0, 59))
        k = rng.choices([1, 2], weights=[75, 25])[0]
        rows.append([
            ts.strftime("%Y/%m/%d %H:%M:%S"),
            age_band(pick_age()),
            labels[s],
            pick_again(s),
            ", ".join(rng.sample(CHANNELS[:4], k)),
            pick_comment(s),
        ])
    path = OUT / "minato_google-form.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["タイムスタンプ", "年代を教えてください", "今日のレッスンの満足度を教えてください",
                    "また参加したいですか？", "このレッスンを何で知りましたか？（複数選択可）", "ご意見・ご感想"])
        w.writerows(rows)
    print(path.relative_to(ROOT), len(rows), "件")


# ----- さくら教室: 手入力の Excel -----
def make_sakura() -> None:
    zenkaku = str.maketrans("12345", "１２３４５")
    again = {"はい": "○", "いいえ": "×", "未定": "△"}
    channel = {"Instagram": "インスタ", "チラシ": "チラシ", "知人の紹介": "友人", "Web検索": "ネット", "その他": "その他"}
    wb = Workbook()
    ws = wb.active
    ws.title = "10月"
    ws["A1"] = "10月 体験レッスン アンケート（さくら教室）"
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["日付", "年齢", "満足度(1〜5)", "再参加", "きっかけ", "感想"])
    n = 0
    for d in lesson_days(38):
        s = pick_satisfaction()
        sat: str | int = s
        if rng.random() < 0.2:
            sat = str(s).translate(zenkaku)  # 全角で入力されたもの
        ws.append([d, pick_age(), sat, again[pick_again(s)], channel[rng.choice(CHANNELS)], pick_comment(s)])
        n += 1
    ws.append([date(2026, 10, 18), 43, "", "○", "インスタ", "とても楽しかったです。"])  # 満足度の記入漏れ
    n += 1
    for row in ws.iter_rows(min_row=4, max_col=1):
        row[0].number_format = "yyyy/m/d"
    ws.column_dimensions["F"].width = 40
    path = OUT / "sakura_handwritten.xlsx"
    wb.save(path)
    print(path.relative_to(ROOT), n, "件")


# ----- ひがし教室: 紙のアンケートを打ち込んだ CSV（Shift_JIS） -----
def make_higashi() -> None:
    words = {5: "満足", 4: "やや満足", 3: "普通", 2: "やや不満", 1: "不満"}
    again = {"はい": "したい", "いいえ": "しない", "未定": "わからない"}
    channel = {"Instagram": "インスタグラム", "チラシ": "チラシ", "知人の紹介": "知り合い", "Web検索": "検索", "その他": "その他"}
    rows = []
    for d in lesson_days(32):
        s = pick_satisfaction()
        rows.append([f"{d.month}/{d.day}", age_band(pick_age()), words[s], again[pick_again(s)],
                     channel[rng.choice(CHANNELS)], pick_comment(s)])
    rows.insert(9, rows[8][:])  # 同じ回答の二重入力
    rows.insert(15, ["", "", "", "", "", ""])  # 空行
    rows.append(["10/25", "40代", "6", "したい", "チラシ", ""])  # 範囲外の満足度
    rows.append(["10/26", "", "満足", "したい", "", "また来ます"])  # 年代の記入漏れ（「不明」として集計に入れる）
    rows.append(["10/28", "50代", "やや満足", "したい", "チラシ", "次回の日程を電話で教えてください。090-0000-0000"])  # ご意見に電話番号
    path = OUT / "higashi_paper.csv"
    with path.open("w", encoding="cp932", newline="") as f:
        w = csv.writer(f)
        w.writerow(["記入日", "年代", "満足度", "またの参加", "知ったきっかけ", "コメント"])
        w.writerows(rows)
    print(path.relative_to(ROOT), len(rows), "行")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_minato()
    make_sakura()
    make_higashi()
