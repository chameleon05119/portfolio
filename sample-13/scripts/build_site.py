"""作品集に載せる紹介ページ（site/）を組み立てる。

    uv run python scripts/build_site.py

1. サンプルの回答ファイルからレポートを作る（out/）
2. 入力 3 ファイルの先頭の行を表にして、site_src/index.html の目印に差し込む
3. 入力・出力・設定ファイルを site/files/ にコピーする
4. macOS ならクイックルックで Excel のサマリーシートを画像にする（site/files/report-summary.png）
"""

from __future__ import annotations

import contextlib
import csv
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from html import escape
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from survey_report.cli import main as run_report  # noqa: E402
from survey_report.config import load_config  # noqa: E402

SITE = ROOT / "site"
FILES = SITE / "files"
PREVIEW_ROWS = 6


def preview_rows(src) -> tuple[list[str], list[list[str]], str]:
    """（見出し, 先頭の行, 説明）を返す。Excel はタイトル行も見えるように上から読む。"""
    if src.file.suffix == ".csv":
        with src.file.open(encoding=src.encoding, newline="") as f:
            table = list(csv.reader(f))
        note = f"CSV・文字コード {'UTF-8' if 'utf' in src.encoding else 'Shift_JIS'}"
        return table[0], table[1:1 + PREVIEW_ROWS], note
    ws = load_workbook(src.file, data_only=True).worksheets[0]
    table = [["" if v is None else (v.strftime("%Y/%m/%d") if hasattr(v, "strftime") else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
    title = table[0][0]
    return table[src.header_row - 1], table[src.header_row:src.header_row + PREVIEW_ROWS], f"Excel・1 行目にタイトル「{title}」、見出しは {src.header_row} 行目"


def table_html(header: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{escape(h)}</th>" for h in header)
    body = "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in r) + "</tr>" for r in rows)
    return f'<div class="scroll" tabindex="0"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>'


def main() -> None:
    os.chdir(ROOT)
    out = ROOT / "out"
    if run_report(["--config", str(ROOT / "config.toml"), "--out", str(out)]) != 0:
        raise SystemExit(1)
    cfg = load_config(ROOT / "config.toml")

    blocks = []
    for src in cfg.sources:
        header, rows, note = preview_rows(src)
        blocks.append(
            f'<figure class="input"><figcaption><b>{escape(src.name)}</b><span>{escape(note)}</span>'
            f'<a href="files/{escape(src.file.name)}" download>ファイル</a></figcaption>{table_html(header, rows)}</figure>'
        )

    html = (ROOT / "site_src" / "index.html").read_text(encoding="utf-8")
    html = html.replace("<!-- inputs -->", "\n".join(blocks))
    # 設定ファイルのうち、最初の教室（[[sources]] 1 つぶん）を見本として載せる
    first_source = "[[sources]]" + (ROOT / "config.toml").read_text(encoding="utf-8").split("[[sources]]")[1]
    html = html.replace("<!-- config -->", escape(first_source.strip()))

    # コマンドの実行結果（ターミナルの表示）とテストの件数は、実際に動かした結果を差し込む（手書きでずれないように）
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        run_report(["--config", "config.toml", "--out", "out"])
    lines = [line for line in buf.getvalue().splitlines() if not line.startswith("書き出し")]
    lines.append("書き出し: out/report.xlsx / out/report.html")
    html = html.replace("<!-- run -->", escape("\n".join(lines)))
    tests = subprocess.run([sys.executable, "-m", "pytest", "-q", "--collect-only"], cwd=ROOT, capture_output=True, text=True).stdout
    count = re.search(r"(\d+) tests? collected", tests)
    html = html.replace("<!-- tests -->", count.group(1) if count else "")

    if SITE.exists():
        shutil.rmtree(SITE)
    shutil.copytree(ROOT / "site_src", SITE, ignore=shutil.ignore_patterns("index.html"))
    (SITE / "index.html").write_text(html, encoding="utf-8")
    FILES.mkdir(parents=True, exist_ok=True)
    for src in cfg.sources:
        shutil.copy(src.file, FILES / src.file.name)
    shutil.copy(ROOT / "config.toml", FILES / "config.toml")
    shutil.copy(out / "report.xlsx", FILES / "report.xlsx")
    shutil.copy(out / "report.html", FILES / "report.html")

    # Excel のサマリーシートの見た目（クイックルックは 1 枚目のシートの表を描く。グラフは描かれない）
    if shutil.which("qlmanage"):
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(["qlmanage", "-t", "-s", "1400", "-o", tmp, str(out / "report.xlsx")], check=True, capture_output=True, timeout=60)
            shutil.copy(Path(tmp) / "report.xlsx.png", FILES / "report-summary.png")
    print(f"組み立て: {SITE.relative_to(ROOT)}/")


if __name__ == "__main__":
    main()
