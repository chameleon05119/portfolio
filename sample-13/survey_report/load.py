"""回答ファイル（CSV・Excel）を「見出し → 値」の行にして読む。"""

from __future__ import annotations

import csv
import zipfile
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import Source


@dataclass(frozen=True)
class RawRow:
    source: str  # 教室名
    file: str  # ファイル名
    line: int  # ファイル上の行番号（Excel の行番号・テキストエディタで見える CSV の行番号）
    values: dict[str, Any]  # レポートの項目 → 元の値


class LoadError(Exception):
    pass


def read_source(src: Source) -> list[RawRow]:
    if not src.file.exists():
        raise LoadError(f"{src.name}: ファイルが見つかりません: {src.file}")
    suffix = src.file.suffix.lower()
    if suffix == ".csv":
        table = _read_csv(src)
    elif suffix in (".xlsx", ".xlsm"):
        table = _read_xlsx(src)
    else:
        raise LoadError(f"{src.name}: 読めない形式です（.csv / .xlsx に対応）: {src.file.name}")

    # 見出しの行（header_row 番目の行）を探す
    header_at = next((i for i, (line, _) in enumerate(table) if line >= src.header_row), None)
    if header_at is None:
        raise LoadError(f"{src.name}: 見出しの行（{src.header_row} 行目）がありません（ファイルが空か、行が足りません）")
    header = [str(h).strip() if h is not None else "" for h in table[header_at][1]]
    positions = {}
    for field, title in src.columns.items():
        if title not in header:
            raise LoadError(f"{src.name}: 見出し「{title}」が見つかりません（{src.header_row} 行目: {' / '.join(h for h in header if h)}）")
        positions[field] = header.index(title)

    rows = []
    for line, cells in table[header_at + 1:]:
        values = {f: (cells[p] if p < len(cells) else None) for f, p in positions.items()}
        rows.append(RawRow(src.name, src.file.name, line, values))
    return rows


def _read_csv(src: Source) -> list[tuple[int, list[Any]]]:
    """(その記録が始まる行番号, セル) の並び。ご意見に改行が入っていても、ファイル上の行番号がずれないようにする。"""
    try:
        with src.file.open(encoding=src.encoding, newline="") as f:
            reader = csv.reader(f)
            out = []
            start = 1
            for cells in reader:
                out.append((start, cells))
                start = reader.line_num + 1
            return out
    except UnicodeDecodeError:
        raise LoadError(f"{src.name}: 文字コード {src.encoding} で読めません。config.toml の encoding を確かめてください（Shift_JIS なら cp932）") from None
    except LookupError:
        raise LoadError(f"{src.name}: encoding「{src.encoding}」という文字コードはありません（utf-8-sig / cp932 など）") from None


def _read_xlsx(src: Source) -> list[tuple[int, list[Any]]]:
    try:
        wb = load_workbook(src.file, read_only=True, data_only=True)  # 数式は計算結果の値で読む
    except (zipfile.BadZipFile, InvalidFileException, KeyError):
        raise LoadError(f"{src.name}: Excel ファイルとして開けません（壊れているか、.xlsx ではない形式です）: {src.file.name}") from None
    ws = wb.worksheets[0]
    table = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
    wb.close()
    return table
