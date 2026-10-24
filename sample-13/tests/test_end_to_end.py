"""サンプルの 3 ファイルからレポートを作り、件数と Excel・HTML の中身を確かめる。エラー時の表示も確かめる。"""

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from survey_report.cli import main

ROOT = Path(__file__).resolve().parent.parent


def test_sample_report(tmp_path, capsys):
    assert main(["--config", str(ROOT / "config.toml"), "--out", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "全体" in out and "回答 119 件" in out
    assert "集計から外した行 2 件・要確認 4 件・空行 1 行" in out

    wb = load_workbook(tmp_path / "report.xlsx")
    assert wb.sheetnames == ["サマリー", "回答一覧", "自由記述", "要確認"]

    summary = wb["サマリー"]
    venues = {summary.cell(r, 1).value: summary.cell(r, 2).value for r in range(7, 11)}
    assert venues == {"みなと教室": 46, "さくら教室": 38, "ひがし教室": 35, "全体": 119}
    assert len(summary._charts) == 3

    # 満足度の内訳（5〜1）の合計が回答数と一致する
    for r in range(7, 11):
        assert sum(summary.cell(r, c).value for c in range(4, 9)) == summary.cell(r, 2).value

    assert wb["回答一覧"].max_row == 119 + 1
    issues = wb["要確認"]
    rows = [[c.value for c in r] for r in issues.iter_rows(min_row=2)]
    assert [r[0] for r in rows] == ["集計から外した", "集計から外した", "集計に入れた", "集計に入れた"]
    reasons = " ".join(r[4] for r in rows)
    for word in ("満足度が空欄", "範囲外", "二重入力の疑い", "電話番号"):
        assert word in reasons

    # 行番号は、元のファイルで見える行番号（さくら教室は見出しが 3 行目なので、データは 4 行目から）
    sakura = next(r for r in rows if r[1] == "さくら教室")
    assert sakura[3] == 4 + 38  # 38 件のあとの行

    html = (tmp_path / "report.html").read_text(encoding="utf-8")
    assert "119 件" in html and "要確認" in html
    assert "090-0000-0000" not in html  # 電話番号は伏せる


def write_csv(path: Path, rows: list[list[str]], encoding: str = "utf-8-sig") -> None:
    with path.open("w", encoding=encoding, newline="") as f:
        csv.writer(f).writerows(rows)


def write_config(tmp_path: Path, file: str, extra: str = "", header_row: str = "1", encoding: str = "utf-8-sig") -> Path:
    cfg = tmp_path / "config.toml"
    cfg.write_text(f"""[report]
title = "テスト"
period = "2026年10月"
[[sources]]
name = "テスト教室"
file = "{file}"
encoding = "{encoding}"
header_row = {header_row}
[sources.columns]
satisfaction = "満足度"
comment = "感想"
{extra}
""", encoding="utf-8")
    return cfg


def run(tmp_path: Path, cfg: Path, capsys) -> tuple[int, str, str]:
    code = main(["--config", str(cfg), "--out", str(tmp_path / "out")])
    out, err = capsys.readouterr()
    return code, out, err


def test_comment_that_looks_like_a_formula_stays_text(tmp_path, capsys):
    write_csv(tmp_path / "a.csv", [["満足度", "感想"], ["5", "=1+1"], ["4", "@SUM(A1)"], ["3", "-よかった"]])
    code, _, _ = run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    assert code == 0
    ws = load_workbook(tmp_path / "out" / "report.xlsx")["回答一覧"]
    cells = [ws.cell(r, 7) for r in range(2, 5)]
    assert [c.value for c in cells] == ["=1+1", "@SUM(A1)", "-よかった"]
    assert all(c.data_type == "s" for c in cells)


def test_html_escapes_comments(tmp_path, capsys):
    write_csv(tmp_path / "a.csv", [["満足度", "感想"], ["2", "<script>alert(1)</script>"]])
    run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    html = (tmp_path / "out" / "report.html").read_text(encoding="utf-8")
    assert "<script>alert" not in html and "&lt;script&gt;" in html


def test_reads_xlsx_with_title_rows(tmp_path, capsys):
    wb = Workbook()
    ws = wb.active
    ws.append(["タイトル"])
    ws.append([])
    ws.append(["満足度", "感想"])
    ws.append(["５", "楽しかった"])
    ws.append([3, ""])
    wb.save(tmp_path / "a.xlsx")
    code, out, _ = run(tmp_path, write_config(tmp_path, "a.xlsx", header_row="3"), capsys)
    assert code == 0 and "回答   2 件　満足度 4.0" in out


def test_multiline_comment_keeps_file_line_numbers(tmp_path, capsys):
    write_csv(tmp_path / "a.csv", [["満足度", "感想"], ["5", "1行目\n2行目\n3行目"], ["9", "範囲外"]])
    run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    ws = load_workbook(tmp_path / "out" / "report.xlsx")["要確認"]
    assert ws.cell(2, 4).value == 5  # 見出し 1 行 + 3 行のご意見のあと = ファイル上は 5 行目


def test_header_only_file_gives_empty_report(tmp_path, capsys):
    write_csv(tmp_path / "a.csv", [["満足度", "感想"]])
    code, out, _ = run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    assert code == 0 and "回答   0 件　満足度 —" in out


@pytest.mark.parametrize("setup, message", [
    (lambda p: write_csv(p / "a.csv", [["満足度", "感想欄"], ["5", ""]]), "見出し「感想」が見つかりません"),
    (lambda p: write_csv(p / "a.csv", [["満足度", "感想"], ["満足", ""]], encoding="cp932"), "文字コード utf-8-sig で読めません"),
    (lambda p: (p / "a.csv").write_text("", encoding="utf-8"), "見出しの行（1 行目）がありません"),
])
def test_input_errors_are_reported_in_japanese(tmp_path, capsys, setup, message):
    setup(tmp_path)
    code, _, err = run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    assert code == 1 and message in err and "Traceback" not in err


def test_broken_xlsx_is_reported(tmp_path, capsys):
    (tmp_path / "a.xlsx").write_bytes(b"not a zip")
    code, _, err = run(tmp_path, write_config(tmp_path, "a.xlsx"), capsys)
    assert code == 1 and "Excel ファイルとして開けません" in err


@pytest.mark.parametrize("body, message", [
    ("[report\n", "設定ファイルの書き方に誤りがあります"),
    ('[[sources]]\nname = "全体"\nfile = "a.csv"\n[sources.columns]\nsatisfaction = "満足度"\n', "集計用の名前（全体）と同じです"),
    ('[[sources]]\nname = "A"\nfile = "a.csv"\nheader_row = "3"\n[sources.columns]\nsatisfaction = "満足度"\n', "header_row は 1 以上の数字"),
])
def test_config_errors_are_reported_in_japanese(tmp_path, capsys, body, message):
    cfg = tmp_path / "config.toml"
    cfg.write_text(body, encoding="utf-8")
    code, _, err = run(tmp_path, cfg, capsys)
    assert code == 1 and message in err


def test_unwritable_output_is_reported(tmp_path, capsys, monkeypatch):
    write_csv(tmp_path / "a.csv", [["満足度", "感想"], ["5", ""]])

    def deny(*args, **kwargs):
        raise PermissionError

    monkeypatch.setattr("survey_report.cli.write_xlsx", deny)
    code, _, err = run(tmp_path, write_config(tmp_path, "a.csv"), capsys)
    assert code == 1 and "Excel で開いている場合は閉じて" in err
